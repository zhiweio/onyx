"""Word report templates: validation, storage, projection, fork, and rendering.

The uploaded file is the layout reference. Tokens inside the document are
not extracted or stored.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Generator

import pytest
from docx import Document
from sqlalchemy import delete
from sqlalchemy.orm import Session

from onyx.db.enums import (
    ReportTemplateKind,
    SystemCatalogCategory,
)
from onyx.db.models import ReportTemplate, SystemReportTemplate, User
from onyx.db.report_template import (
    attach_docx_asset,
    create_report_template,
    delete_report_template,
    read_docx_asset,
)
from onyx.db.system_catalog.fork import fork_system_report_template_for_user
from onyx.db.system_catalog.publish import (
    publish_system_report_template,
    unpublish_system_report_template,
)
from onyx.db.system_catalog.report_template import (
    attach_catalog_docx_asset,
    create_system_report_template,
)
from onyx.error_handling.error_codes import OnyxErrorCode
from onyx.error_handling.exceptions import OnyxError
from onyx.file_store.file_store import get_default_file_store
from onyx.report_templates.docx_template import validate_docx_asset
from onyx.server.features.scenario.runtime import render_report_template_section
from tests.external_dependency_unit.conftest import create_test_user


def build_docx() -> bytes:
    document = Document()
    document.add_heading("月度关账报告", 0)
    document.add_paragraph("主体：某某股份有限公司  期间：2026年8月")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def template_user(db_session: Session) -> User:
    return create_test_user(db_session, "docx_tpl")


@pytest.fixture
def unique_slug(request: pytest.FixtureRequest) -> str:
    return f"docx_{abs(hash(request.node.name)) % 10**8}"


# ── validation ──────────────────────────────────────────────────────────────


def test_validation_accepts_a_readable_docx() -> None:
    validate_docx_asset(build_docx())


def test_validation_rejects_a_file_that_is_not_a_docx() -> None:
    with pytest.raises(OnyxError) as caught:
        validate_docx_asset(b"definitely not a zip")
    assert caught.value.error_code is OnyxErrorCode.INVALID_INPUT


def test_validation_rejects_a_zip_without_a_word_document() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hello")
    with pytest.raises(OnyxError) as caught:
        validate_docx_asset(buffer.getvalue())
    assert caught.value.error_code is OnyxErrorCode.INVALID_INPUT


def test_validation_refuses_declared_xml_entities() -> None:
    """An uploaded document is untrusted input."""
    bomb = (
        b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "boom">]>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&a;</w:t>'
        b"</w:r></w:p></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", bomb)
    with pytest.raises(OnyxError) as caught:
        validate_docx_asset(buffer.getvalue())
    assert caught.value.error_code is OnyxErrorCode.INVALID_INPUT


# ── user templates ──────────────────────────────────────────────────────────


@pytest.fixture
def user_template(
    db_session: Session, template_user: User, unique_slug: str
) -> Generator[ReportTemplate, None, None]:
    template = create_report_template(
        db_session,
        user=template_user,
        name="Docx template under test",
        slug=unique_slug,
        description="",
        body="# Outline\n",
    )
    yield template
    _purge_template(db_session, template.slug)


def _purge_template(db_session: Session, slug: str) -> None:
    row = db_session.scalar(
        delete(ReportTemplate)
        .where(ReportTemplate.slug == slug)
        .returning(ReportTemplate.asset_file_id)
    )
    db_session.commit()
    if row:
        get_default_file_store().delete_file(row, error_on_missing=False)


def test_attaching_a_docx_switches_kind_and_stores_the_file(
    db_session: Session, template_user: User, user_template: ReportTemplate
) -> None:
    assert user_template.kind is ReportTemplateKind.MARKDOWN
    asset_bytes = build_docx()

    updated = attach_docx_asset(
        db_session,
        user_template,
        template_user,
        asset_bytes=asset_bytes,
        filename="close.docx",
    )

    assert updated.kind is ReportTemplateKind.DOCX
    assert updated.asset_filename == "close.docx"
    assert updated.asset_file_id is not None
    assert updated.asset_sha256 is not None
    assert read_docx_asset(updated) == asset_bytes


def test_replacing_the_asset_drops_the_previous_blob(
    db_session: Session, template_user: User, user_template: ReportTemplate
) -> None:
    first = attach_docx_asset(
        db_session,
        user_template,
        template_user,
        asset_bytes=build_docx(),
        filename="v1.docx",
    )
    first_file_id = first.asset_file_id
    assert first_file_id is not None

    document = Document()
    document.add_paragraph("Replacement layout.")
    buffer = io.BytesIO()
    document.save(buffer)
    second = attach_docx_asset(
        db_session,
        user_template,
        template_user,
        asset_bytes=buffer.getvalue(),
        filename="v2.docx",
    )

    assert second.asset_file_id != first_file_id
    assert second.asset_filename == "v2.docx"


def test_reading_the_asset_of_a_markdown_template_is_not_found(
    user_template: ReportTemplate,
) -> None:
    with pytest.raises(OnyxError) as caught:
        read_docx_asset(user_template)
    assert caught.value.error_code is OnyxErrorCode.NOT_FOUND


# ── catalog projection and fork ─────────────────────────────────────────────


@pytest.fixture
def catalog_entry(
    db_session: Session, unique_slug: str
) -> Generator[SystemReportTemplate, None, None]:
    entry = create_system_report_template(
        db_session,
        slug=f"{unique_slug}_cat",
        name="Catalog Word template",
        description="A Word template under test.",
        body="# Guidance\n",
        category=SystemCatalogCategory.TAX,
        tags=[],
    )
    db_session.commit()
    yield entry
    db_session.execute(
        delete(ReportTemplate).where(
            ReportTemplate.system_report_template_id == entry.id
        )
    )
    db_session.execute(
        delete(SystemReportTemplate).where(SystemReportTemplate.id == entry.id)
    )
    db_session.commit()


def test_publish_projects_the_word_asset(
    db_session: Session, template_user: User, catalog_entry: SystemReportTemplate
) -> None:
    attach_catalog_docx_asset(
        db_session, catalog_entry, asset_bytes=build_docx(), filename="cat.docx"
    )
    db_session.commit()

    projection = publish_system_report_template(
        db_session, catalog_entry, publisher=template_user
    )
    db_session.commit()

    assert projection.kind is ReportTemplateKind.DOCX
    assert projection.asset_filename == "cat.docx"
    # The projection borrows the catalog's blob rather than duplicating it.
    assert projection.asset_file_id == catalog_entry.asset_file_id


def test_unpublish_keeps_the_catalog_asset_readable(
    db_session: Session, template_user: User, catalog_entry: SystemReportTemplate
) -> None:
    """The projection only borrows the blob, so dropping it must not delete it."""
    asset_bytes = build_docx()
    attach_catalog_docx_asset(
        db_session, catalog_entry, asset_bytes=asset_bytes, filename="cat.docx"
    )
    db_session.commit()
    publish_system_report_template(db_session, catalog_entry, publisher=template_user)
    db_session.commit()

    unpublish_system_report_template(db_session, catalog_entry)
    db_session.commit()

    assert catalog_entry.asset_file_id is not None
    payload = get_default_file_store().read_file(catalog_entry.asset_file_id).read()
    assert payload == asset_bytes


def test_fork_copies_the_asset_so_the_user_owns_it(
    db_session: Session, template_user: User, catalog_entry: SystemReportTemplate
) -> None:
    asset_bytes = build_docx()
    attach_catalog_docx_asset(
        db_session, catalog_entry, asset_bytes=asset_bytes, filename="cat.docx"
    )
    db_session.commit()
    publish_system_report_template(db_session, catalog_entry, publisher=template_user)
    db_session.commit()

    fork = fork_system_report_template_for_user(
        db_session, catalog_entry, template_user
    )
    db_session.commit()

    assert fork.kind is ReportTemplateKind.DOCX
    # A fork owns its bytes: deleting it must not disturb the catalog entry.
    assert fork.asset_file_id != catalog_entry.asset_file_id
    assert read_docx_asset(fork) == asset_bytes

    delete_report_template(db_session, fork, template_user)

    assert catalog_entry.asset_file_id is not None
    assert (
        get_default_file_store().read_file(catalog_entry.asset_file_id).read()
        == asset_bytes
    )


# ── SCENARIO.md rendering ───────────────────────────────────────────────────


def test_markdown_template_renders_its_outline(
    user_template: ReportTemplate,
) -> None:
    rendered = "\n".join(render_report_template_section(user_template))
    assert "# Outline" in rendered
    assert "fill_template.py" not in rendered


def test_docx_template_renders_the_file_as_a_reference(
    db_session: Session, template_user: User, user_template: ReportTemplate
) -> None:
    template = attach_docx_asset(
        db_session,
        user_template,
        template_user,
        asset_bytes=build_docx(),
        filename="close.docx",
    )

    rendered = "\n".join(render_report_template_section(template))

    assert f"/workspace/managed/report_templates/{template.slug}.docx" in rendered
    assert "layout and style reference" in rendered
    assert "fill_template.py" not in rendered
    assert "report-data.json" not in rendered
    assert "Placeholders to fill" not in rendered
    assert "# Outline" in rendered


def test_official_builder_returns_a_readable_docx() -> None:
    from onyx.system_catalog.builtin.word.generate import generate_official_docx

    asset_bytes = generate_official_docx("monthly_close")
    validate_docx_asset(asset_bytes)
    assert asset_bytes.startswith(b"PK")
