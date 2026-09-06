from __future__ import annotations

import importlib.util
import io
from pathlib import Path

import pytest
from docx import Document

_FILL_PATH = (
    Path(__file__).resolve().parents[4]
    / "onyx"
    / "skills"
    / "builtin"
    / "docx"
    / "scripts"
    / "fill_template.py"
)
_SPEC = importlib.util.spec_from_file_location("docx_fill_template", _FILL_PATH)
assert _SPEC is not None and _SPEC.loader is not None
fill_template = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(fill_template)


def _prototype_docx() -> bytes:
    document = Document()
    document.add_paragraph("主体：{{entity_name}}")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "发现"
    table.cell(0, 1).text = "严重程度"
    table.cell(0, 2).text = "金额"
    table.cell(1, 0).text = "{{findings.title}}"
    table.cell(1, 1).text = "{{findings.severity}}"
    table.cell(1, 2).text = "{{findings.amount}}"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_prototype_row_clones_once_per_array_item() -> None:
    document = Document(io.BytesIO(_prototype_docx()))
    result = fill_template.fill_document(
        document,
        {
            "entity_name": "某某公司",
            "findings": [
                {"title": "未达账", "severity": "高", "amount": "120万"},
                {"title": "发票缺失", "severity": "中", "amount": "8万"},
                {"title": "费用错科目", "severity": "低", "amount": "1万"},
            ],
        },
    )
    assert result["leftover"] == []
    assert result["row_clones"] == 2
    texts = [row.cells[0].text for row in document.tables[0].rows]
    assert texts == ["发现", "未达账", "发票缺失", "费用错科目"]


def test_leftover_tokens_are_reported() -> None:
    document = Document(io.BytesIO(_prototype_docx()))
    result = fill_template.fill_document(document, {"entity_name": "某某公司"})
    assert "findings.title" in result["leftover"]


def test_leftover_tokens_make_main_fail(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    template = tmp_path / "template.docx"
    template.write_bytes(_prototype_docx())
    data = tmp_path / "data.json"
    data.write_text('{"entity_name": "某某公司"}', encoding="utf-8")
    output = tmp_path / "out.docx"
    monkeypatch.setattr(
        "sys.argv",
        [
            "fill_template.py",
            "--template",
            str(template),
            "--data",
            str(data),
            "--output",
            str(output),
        ],
    )
    with pytest.raises(ValueError, match="unfilled placeholders"):
        fill_template.main()
    assert not output.exists()


def test_unused_json_keys_are_reported_not_fatal() -> None:
    document = Document(io.BytesIO(_prototype_docx()))
    result = fill_template.fill_document(
        document,
        {
            "entity_name": "某某公司",
            "unused_note": "ignore me",
            "findings": [{"title": "A", "severity": "高", "amount": "1"}],
        },
    )
    assert result["leftover"] == []
    assert "unused_note" in result["unused_keys"]
