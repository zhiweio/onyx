"""Fused financial-report-analysis skill stays parseable and scriptable."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from onyx.db.enums import SystemCatalogCategory
from onyx.skills.built_in import BUILT_IN_SKILLS, BUILTIN_SKILLS_PATH
from onyx.skills.metadata import parse_skill_document
from onyx.system_catalog.builtin.manifest import BUILT_IN_SKILL_ENTRIES

_SKILL_ID = "financial-report-analysis"
_SKILL_DIR = BUILTIN_SKILLS_PATH / _SKILL_ID
_SCRIPTS = _SKILL_DIR / "scripts"
_SELF_CHECK = _SCRIPTS / "self_check_report.py"


def test_financial_report_analysis_is_registered_and_parses() -> None:
    definition = BUILT_IN_SKILLS[_SKILL_ID]
    document = parse_skill_document(
        (definition.source_dir / "SKILL.md").read_bytes(),
        directory_name=_SKILL_ID,
    )
    assert document.metadata.name == _SKILL_ID
    assert document.instructions_markdown
    assert "hithink-a-share" in (document.metadata.optional_mcp or [])
    assert "qcc-company" in (document.metadata.optional_mcp or [])
    assert "company-credit" in (document.metadata.optional_mcp or [])


def test_financial_report_analysis_is_in_the_gallery_manifest() -> None:
    entry = next(item for item in BUILT_IN_SKILL_ENTRIES if item.slug == _SKILL_ID)
    assert entry.built_in_skill_id == _SKILL_ID
    assert entry.category is SystemCatalogCategory.REPORT
    assert entry.name == "财报解读"


def test_analyze_financials_sample_detects_anomalies(tmp_path: Path) -> None:
    sample = json.loads(
        subprocess.check_output(
            [sys.executable, str(_SCRIPTS / "analyze_financials.py"), "--sample"],
            text=True,
        )
    )
    sample_path = tmp_path / "sample.json"
    sample_path.write_text(json.dumps(sample, ensure_ascii=False), encoding="utf-8")
    report = json.loads(
        subprocess.check_output(
            [
                sys.executable,
                str(_SCRIPTS / "analyze_financials.py"),
                str(sample_path),
                "--json",
            ],
            text=True,
        )
    )
    assert report["company"] == sample["company"]
    assert report["summary"]["total_anomalies"] >= 1


def test_normalize_statements_emits_trend_rows(tmp_path: Path) -> None:
    sample = subprocess.check_output(
        [sys.executable, str(_SCRIPTS / "analyze_financials.py"), "--sample"],
        text=True,
    )
    input_path = tmp_path / "statements.json"
    trend_path = tmp_path / "trend.json"
    input_path.write_text(sample, encoding="utf-8")
    summary = json.loads(
        subprocess.check_output(
            [
                sys.executable,
                str(_SCRIPTS / "normalize_statements.py"),
                "--input",
                str(input_path),
                "--trend",
                str(trend_path),
            ],
            text=True,
        )
    )
    assert summary["income_rows"] == 8
    trend = json.loads(trend_path.read_text(encoding="utf-8"))
    assert trend["income"][0]["EndDate"] == "2023-03-31"
    assert trend["income"][0]["OperatingRevenue"] == 50_000_000
    assert trend["cashflow"][-1]["NetOperateCashFlow"] == 2_000_000


def _instructions() -> str:
    definition = BUILT_IN_SKILLS[_SKILL_ID]
    return parse_skill_document(
        (definition.source_dir / "SKILL.md").read_bytes(),
        directory_name=_SKILL_ID,
    ).instructions_markdown


def test_skill_instructions_ban_internal_citations() -> None:
    """Reader-facing citations only: no MCP names, no sandbox temp paths."""
    instructions = _instructions()
    assert "禁止写进报告" in instructions
    assert "/tmp/" in instructions
    assert "同花顺（iFinD）" in instructions
    # The old rule explicitly allowed MCP tool citations.
    assert "MCP 工具或网页 URL" not in instructions


def test_skill_dual_delivery_contract() -> None:
    """HTML goes through slideblocks; Word goes through the docx skill."""
    instructions = _instructions()
    assert "slideblocks" in instructions
    assert "docx 技能" in instructions
    assert "self_check_report.py" in instructions
    # The HTML report draws with slideblocks' own render routes; chart PNGs
    # produced by other skills are Word-only.
    assert "内置渲染路径" in instructions
    assert "报告图表一律 PNG" not in instructions
    assert "图表一律嵌入" not in instructions
    assert (_SKILL_DIR / "references" / "docx-report-template.md").is_file()
    # The hand-rolled HTML skeleton is retired.
    assert not (_SKILL_DIR / "assets" / "report_template.html").exists()
    stale = "\n".join(
        path.read_text(encoding="utf-8")
        for path in (
            _SKILL_DIR / "SKILL.md",
            _SKILL_DIR / "references" / "output_template.md",
        )
    )
    assert "report_template.html" not in stale


def test_skill_requires_eight_charts() -> None:
    assert "至少产出 8 张" in _instructions()


def _run_self_check(target: Path, min_figures: int = 1) -> tuple[int, str]:
    process = subprocess.run(
        [
            sys.executable,
            str(_SELF_CHECK),
            str(target),
            "--min-figures",
            str(min_figures),
        ],
        capture_output=True,
        text=True,
    )
    return process.returncode, process.stdout + process.stderr


def test_self_check_passes_a_clean_html_report(tmp_path: Path) -> None:
    report = tmp_path / "report.html"
    report.write_text(
        "<html><body>"
        "<h1>雪龙集团 财报解读</h1>"
        + "".join('<img src="data:image/png;base64,AAA">' for _ in range(2))
        + "<p>数据来源：公司2025年年度报告第 12 页</p>"
        "<p>免责声明：本解读不构成投资建议。</p>"
        "</body></html>",
        encoding="utf-8",
    )
    exit_code, output = _run_self_check(report)
    assert exit_code == 0, output
    assert "FAIL" not in output


def test_self_check_counts_inline_svg_figures(tmp_path: Path) -> None:
    """slideblocks draws HTML figures as inline SVG carriers, not <img> tags."""
    report = tmp_path / "report.html"
    report.write_text(
        "<html><body>"
        "<h1>雪龙集团 财报解读</h1>"
        + "".join(
            '<svg data-slideblocks-render-route="figure:comparison"></svg>'
            for _ in range(2)
        )
        + "<p>数据来源：公司2025年年度报告第 12 页</p>"
        "<p>免责声明：本解读不构成投资建议。</p>"
        "</body></html>",
        encoding="utf-8",
    )
    exit_code, output = _run_self_check(report)
    assert exit_code == 0, output
    assert "2 figures" in output


def test_self_check_flags_leaky_html_report(tmp_path: Path) -> None:
    report = tmp_path / "report.html"
    report.write_text(
        "<html><body>"
        "<h1>雪龙集团（草稿待修改）</h1>"
        '<img src="outputs/charts/a.png">'
        "<p>数据来自 hithink-a-share（MCP），明细见 /tmp/report_data.csv 与 {{公司名}}</p>"
        "</body></html>",
        encoding="utf-8",
    )
    exit_code, output = _run_self_check(report)
    assert exit_code == 1
    assert "FAIL internal_paths" in output
    assert "FAIL mcp_names" in output
    assert "FAIL draft_markers" in output
    assert "FAIL placeholders" in output


def test_self_check_passes_a_clean_docx_report(tmp_path: Path) -> None:
    from docx import Document

    report = tmp_path / "report.docx"
    document = Document()
    document.add_heading("雪龙集团 财报解读", 1)
    document.add_paragraph("数据来源：公司2025年年度报告第 12 页")
    document.add_paragraph("免责声明：本解读不构成投资建议。")
    document.save(str(report))
    exit_code, output = _run_self_check(report, min_figures=0)
    assert exit_code == 0, output


def test_self_check_flags_a_leaky_docx_report(tmp_path: Path) -> None:
    from docx import Document

    report = tmp_path / "report.docx"
    document = Document()
    document.add_heading("雪龙集团 财报解读", 1)
    document.add_paragraph("原始 JSON 存放在 outputs/mcp/hithink/，TODO：补充同业数据")
    document.save(str(report))
    exit_code, output = _run_self_check(report, min_figures=0)
    assert exit_code == 1
    assert "FAIL internal_paths" in output
    assert "FAIL draft_markers" in output
