"""Convert Markdown to a DOCX document using the shared GFM AST + python-docx.

Used by the build session "export as DOCX" feature. ``parse_markdown`` builds
the same mistune tree as PDF export; ``python-docx`` writes OOXML. Both are
pure-Python, so the conversion needs no external binary.

Supported constructs (covering what the LLM-generated documents emit):
headings, bold/italic/strikethrough/inline-code, bulleted/numbered/nested
lists (with loose-list continuation paragraphs and preserved ordered-list
start values), blockquotes, fenced code blocks, GFM tables, hyperlinks
(carrying inherited inline formatting), images (embedded when an
``image_loader`` supplies bytes; otherwise alt text), inline
``<br>`` line breaks, HTML entities, and horizontal rules. Other raw HTML is
dropped rather than shown as literal markup.
"""

from dataclasses import dataclass, replace
from html import unescape
from io import BytesIO
from typing import Any, cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.package import OpcPackage
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor, Twips
from docx.styles.style import ParagraphStyle
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from onyx.server.features.build.session.md_document import (
    Node,
    parse_markdown,
)
from onyx.server.features.build.session.md_export_style import (
    ACCENT,
    BLOCK_INDENT_IN,
    BLOCK_SPACE_PT,
    BODY_EAST_ASIA,
    BODY_FONT,
    BODY_SIZE_PT,
    BODY_SPACE_AFTER_PT,
    CELL_SIZE_PT,
    CODE_BG,
    CODE_SIZE_PT,
    COMPACT_SPACE_PT,
    DOC_GRID_LINE_PITCH,
    FOOTER_DISTANCE_MM,
    HEADER_DISTANCE_MM,
    HEADING_EAST_ASIA,
    HEADING_FONT,
    HEADING_SIZES_PT,
    HEADING_SPACE_AFTER_PT,
    HEADING_SPACE_BEFORE_PT,
    INK,
    LINK,
    MARGIN_BOTTOM_MM,
    MARGIN_LEFT_MM,
    MARGIN_RIGHT_MM,
    MARGIN_TOP_MM,
    MONO_FONT,
    MUTED,
    PAGE_HEIGHT_MM,
    PAGE_WIDTH_MM,
    RULE,
    TABLE_ALT_FILL,
    TABLE_BORDER,
    TABLE_HEADER_FILL,
    WHITE,
    WORD_BODY_LINE_PT,
    WORD_CELL_LINE_PT,
    WORD_CODE_LINE_PT,
    WORD_COMPACT_LINE_PT,
    first_heading_text,
    word_heading_line_pt,
)
from onyx.server.features.build.session.md_images import (
    ImageLoader,
    attach_image_bytes,
    fit_image_display_size,
    image_bytes,
)

_MONOSPACE_FONT = MONO_FONT
_CODE_FONT_SIZE = Pt(CODE_SIZE_PT)
# Links use a "Hyperlink" character style: muted blue, no underline.
_STYLE_HYPERLINK = "Hyperlink"
_STYLE_ID_HYPERLINK = "Hyperlink"


def _rgb(hex6: str) -> RGBColor:
    return RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))


_LINK_COLOR = _rgb(LINK)
_HEADER_COLOR = _rgb(WHITE)
_HEADING_COLOR = _rgb(ACCENT)
_MUTED_COLOR = _rgb(MUTED)
# python-docx ships built-in "List Bullet"/"List Number" styles plus numbered
# variants up to level 3 ("List Bullet 2", "List Bullet 3", ...). Deeper nesting
# reuses the level-3 style.
_MAX_LIST_LEVEL = 3
# pandoc indents each list level by 0.5" (720 twips) with a 0.25" hanging marker;
# python-docx's built-in list numbering indents only half as far.
_LIST_INDENT_PER_LEVEL = 720
_LIST_HANGING_INDENT = 360

# Paragraph style names stay aligned with the previous pandoc mapping so
# existing assignment rules (First Paragraph / Body Text / Compact) still hold.
_STYLE_BODY = "Body Text"
_STYLE_FIRST_PARAGRAPH = "First Paragraph"
_STYLE_COMPACT = "Compact"
_STYLE_IMAGE_CAPTION = "Image Caption"
_STYLE_BLOCK_TEXT = "Block Text"

_BODY_SPACE_AFTER = Pt(BODY_SPACE_AFTER_PT)
_COMPACT_SPACE = Pt(COMPACT_SPACE_PT)
_BLOCK_TEXT_SPACE = Pt(BLOCK_SPACE_PT)
_BLOCK_TEXT_INDENT = Inches(BLOCK_INDENT_IN)

# Footnotes are written as a real Word footnotes part (python-docx has no native
# API for them), so [^n] citations become superscript references that Word links
# to page-bottom notes. Style ids are the spaceless form of the style names.
_STYLE_FOOTNOTE_TEXT = "Footnote Text"
_STYLE_FOOTNOTE_REFERENCE = "Footnote Reference"
_STYLE_ID_FOOTNOTE_TEXT = "FootnoteText"
_STYLE_ID_FOOTNOTE_REFERENCE = "FootnoteReference"
_FOOTNOTES_PARTNAME = "/word/footnotes.xml"
_FOOTNOTES_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)
_FOOTNOTES_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
)
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


@dataclass(frozen=True)
class _Fmt:
    """Inline formatting flags carried down through nested inline nodes."""

    bold: bool = False
    italic: bool = False
    strike: bool = False
    code: bool = False
    color: RGBColor | None = None


def markdown_to_docx_bytes(
    md_text: str,
    *,
    image_loader: ImageLoader | None = None,
) -> bytes:
    """Render Markdown text to the bytes of a .docx file."""
    nodes = parse_markdown(md_text)
    attach_image_bytes(nodes, image_loader)

    document = Document()
    _drop_template_empty_paragraph(document)
    _apply_report_styles(document)
    _apply_cjk_document_defaults(document)
    _apply_page_chrome(document, first_heading_text(nodes))
    footnote_block = next(
        (node for node in nodes if node.get("type") == "footnotes"), None
    )
    footnotes = (
        _Footnotes(document, footnote_block.get("children", []))
        if footnote_block is not None
        else None
    )
    _render_blocks(document, nodes, footnotes)
    if document.element.body.find(qn("w:p")) is None:
        _add_styled_paragraph(document, _STYLE_BODY)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _apply_report_styles(document: DocxDocument) -> None:
    """Configure A4 page, compact body, and navy heading hierarchy."""
    for section in document.sections:
        section.page_width = Mm(PAGE_WIDTH_MM)
        section.page_height = Mm(PAGE_HEIGHT_MM)
        section.left_margin = Mm(MARGIN_LEFT_MM)
        section.right_margin = Mm(MARGIN_RIGHT_MM)
        section.top_margin = Mm(MARGIN_TOP_MM)
        section.bottom_margin = Mm(MARGIN_BOTTOM_MM)
        section.header_distance = Mm(HEADER_DISTANCE_MM)
        section.footer_distance = Mm(FOOTER_DISTANCE_MM)

    styles = document.styles
    existing = {style.name for style in styles}

    def set_east_asia(style: ParagraphStyle, latin: str, east_asia: str) -> None:
        style.font.name = latin
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        r_fonts.set(qn("w:ascii"), latin)
        r_fonts.set(qn("w:hAnsi"), latin)
        r_fonts.set(qn("w:eastAsia"), east_asia)

    normal = cast(ParagraphStyle, styles["Normal"])
    set_east_asia(normal, BODY_FONT, BODY_EAST_ASIA)
    normal.font.size = Pt(BODY_SIZE_PT)
    normal.font.color.rgb = _rgb(INK)
    _set_exact_line_spacing(normal.paragraph_format, WORD_BODY_LINE_PT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = _BODY_SPACE_AFTER

    def ensure(name: str, base: str) -> ParagraphStyle:
        if name not in existing:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles[base]
            existing.add(name)
        return cast(ParagraphStyle, styles[name])

    body = cast(ParagraphStyle, styles[_STYLE_BODY])  # ships in the default template
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = _BODY_SPACE_AFTER
    _set_exact_line_spacing(body.paragraph_format, WORD_BODY_LINE_PT)

    ensure(_STYLE_FIRST_PARAGRAPH, _STYLE_BODY)

    compact = ensure(_STYLE_COMPACT, _STYLE_BODY)
    compact.paragraph_format.space_before = _COMPACT_SPACE
    compact.paragraph_format.space_after = _COMPACT_SPACE
    _set_exact_line_spacing(compact.paragraph_format, WORD_COMPACT_LINE_PT)

    block_text = ensure(_STYLE_BLOCK_TEXT, _STYLE_BODY)
    block_text.paragraph_format.space_before = _BLOCK_TEXT_SPACE
    block_text.paragraph_format.space_after = _BLOCK_TEXT_SPACE
    block_text.paragraph_format.left_indent = _BLOCK_TEXT_INDENT
    block_text.paragraph_format.right_indent = Pt(0)

    caption = ensure(_STYLE_IMAGE_CAPTION, "Caption")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.font.size = Pt(9)
    caption.font.color.rgb = _MUTED_COLOR
    _set_exact_line_spacing(caption.paragraph_format, WORD_COMPACT_LINE_PT)

    ensure(_STYLE_FOOTNOTE_TEXT, "Normal")
    if _STYLE_FOOTNOTE_REFERENCE not in existing:
        reference = styles.add_style(_STYLE_FOOTNOTE_REFERENCE, WD_STYLE_TYPE.CHARACTER)
        reference.font.superscript = True

    if _STYLE_HYPERLINK not in existing:
        hyperlink = styles.add_style(_STYLE_HYPERLINK, WD_STYLE_TYPE.CHARACTER)
        hyperlink.font.color.rgb = _LINK_COLOR

    for level, size in HEADING_SIZES_PT.items():
        heading = styles[f"Heading {level}"]
        set_east_asia(heading, HEADING_FONT, HEADING_EAST_ASIA)
        heading.font.size = Pt(size)
        heading.font.color.rgb = _HEADING_COLOR
        heading.font.bold = True
        _set_exact_line_spacing(
            heading.paragraph_format, word_heading_line_pt(level)
        )
        heading.paragraph_format.space_before = Pt(HEADING_SPACE_BEFORE_PT[level])
        heading.paragraph_format.space_after = Pt(HEADING_SPACE_AFTER_PT[level])
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.keep_together = True
        heading.paragraph_format.widow_control = True

    set_east_asia(body, BODY_FONT, BODY_EAST_ASIA)
    list_styles = (
        "List Bullet",
        "List Number",
        "List Continue",
        "List Bullet 2",
        "List Number 2",
        "List Continue 2",
        "List Bullet 3",
        "List Number 3",
        "List Continue 3",
    )
    for inherited in (
        _STYLE_FIRST_PARAGRAPH,
        _STYLE_COMPACT,
        _STYLE_BLOCK_TEXT,
        _STYLE_IMAGE_CAPTION,
        _STYLE_FOOTNOTE_TEXT,
        *list_styles,
    ):
        if inherited not in existing:
            continue
        style = cast(ParagraphStyle, styles[inherited])
        set_east_asia(style, BODY_FONT, BODY_EAST_ASIA)
        if inherited in list_styles:
            _set_exact_line_spacing(style.paragraph_format, WORD_COMPACT_LINE_PT)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = _COMPACT_SPACE


def _apply_cjk_document_defaults(document: DocxDocument) -> None:
    """Set Word document language and theme fonts the way Word/WPS do.

    Style-level w:eastAsia is not enough: the template defaults to ja-JP
    theme language and theme-linked East-Asian fonts. Document defaults
    plus the theme font scheme make 微软雅黑/黑体 apply to unstyled runs too.
    """
    theme_lang = document.settings.element.find(qn("w:themeFontLang"))
    if theme_lang is None:
        theme_lang = OxmlElement("w:themeFontLang")
        document.settings.element.append(theme_lang)
    theme_lang.set(qn("w:val"), "en-US")
    theme_lang.set(qn("w:eastAsia"), "zh-CN")

    styles_el = document.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, doc_defaults)
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        doc_defaults.append(rpr_default)
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_fonts.set(qn("w:eastAsia"), BODY_EAST_ASIA)
    r_fonts.set(qn("w:cs"), BODY_FONT)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    half_points = str(int(BODY_SIZE_PT * 2))
    for tag in ("sz", "szCs"):
        size_el = rpr.find(qn(f"w:{tag}"))
        if size_el is None:
            size_el = OxmlElement(f"w:{tag}")
            rpr.append(size_el)
        size_el.set(qn("w:val"), half_points)

    ppr_default = doc_defaults.find(qn("w:pPrDefault"))
    if ppr_default is None:
        ppr_default = OxmlElement("w:pPrDefault")
        doc_defaults.append(ppr_default)
    ppr = ppr_default.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        ppr_default.append(ppr)
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), str(int(BODY_SPACE_AFTER_PT * 20)))
    spacing.set(qn("w:line"), str(int(WORD_BODY_LINE_PT * 20)))
    spacing.set(qn("w:lineRule"), "exact")

    _apply_cjk_document_grid(document)

    for rel in document.part.rels.values():
        if rel.reltype != RELATIONSHIP_TYPE.THEME:
            continue
        root = etree.fromstring(rel.target_part.blob)
        for tag, typeface in (
            ("majorFont", HEADING_EAST_ASIA),
            ("minorFont", BODY_EAST_ASIA),
        ):
            for node in root.findall(f".//{{{_A_NS}}}{tag}"):
                east_asia = node.find(f"{{{_A_NS}}}ea")
                if east_asia is not None:
                    east_asia.set("typeface", typeface)
        # Theme parts are generic OPC blobs, not XmlPart.
        rel.target_part._blob = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )


def _apply_cjk_document_grid(document: DocxDocument) -> None:
    """Keep a line-pitch hint without snapping paragraphs to the grid.

    ``type="lines"`` pads each paragraph up to the next grid row. With 1.15
    leading plus space-after, Word then uses two rows per body paragraph.
    """
    for section in document.sections:
        sect_pr = section._sectPr
        doc_grid = sect_pr.find(qn("w:docGrid"))
        if doc_grid is None:
            doc_grid = OxmlElement("w:docGrid")
            sect_pr.append(doc_grid)
        doc_grid.set(qn("w:type"), "default")
        doc_grid.set(qn("w:linePitch"), DOC_GRID_LINE_PITCH)


def _apply_page_chrome(document: DocxDocument, title: str) -> None:
    """Add a running header after page 1 and a centred page number."""
    section = document.sections[0]
    section.different_first_page_header_footer = True
    _fill_header(section.first_page_header.paragraphs[0], "")
    _fill_header(section.header.paragraphs[0], title)
    _fill_footer(section.first_page_footer.paragraphs[0])
    _fill_footer(section.footer.paragraphs[0])


def _fill_header(paragraph: Paragraph, title: str) -> None:
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if not title:
        return
    run = paragraph.add_run(title[:40])
    _set_run_typefaces(run, BODY_FONT, BODY_EAST_ASIA)
    run.font.size = Pt(8)
    run.font.color.rgb = _MUTED_COLOR
    _add_paragraph_border(paragraph, "bottom", RULE)


def _fill_footer(paragraph: Paragraph) -> None:
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_paragraph_border(paragraph, "top", RULE)
    prefix = paragraph.add_run("— ")
    _set_run_typefaces(prefix, BODY_FONT, BODY_EAST_ASIA)
    prefix.font.size = Pt(8)
    prefix.font.color.rgb = _MUTED_COLOR
    _add_page_field(paragraph)
    suffix = paragraph.add_run(" —")
    _set_run_typefaces(suffix, BODY_FONT, BODY_EAST_ASIA)
    suffix.font.size = Pt(8)
    suffix.font.color.rgb = _MUTED_COLOR


def _set_run_typefaces(run: Run, latin: str, east_asia: str) -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _add_paragraph_border(paragraph: Paragraph, edge: str, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    line = OxmlElement(f"w:{edge}")
    line.set(qn("w:val"), "single")
    line.set(qn("w:sz"), "6")
    line.set(qn("w:space"), "4")
    line.set(qn("w:color"), color)
    borders.append(line)


def _add_page_field(paragraph: Paragraph) -> None:
    run = paragraph.add_run()
    _set_run_typefaces(run, BODY_FONT, BODY_EAST_ASIA)
    run.font.size = Pt(8)
    run.font.color.rgb = _MUTED_COLOR
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# --------------------------------------------------------------------------- #
# Footnotes
# --------------------------------------------------------------------------- #
class _PartParent:
    """Minimal parent so a Paragraph built in the footnotes part resolves
    ``paragraph.part`` (used for hyperlink relationships) to that part."""

    def __init__(self, part: XmlPart) -> None:
        self.part = part


class _Footnotes:
    """Builds a Word footnotes part and wires up references to it.

    Word stores footnotes in a separate ``word/footnotes.xml`` part with two
    reserved entries (separator + continuation separator) followed by the real
    notes; the body points at one via ``<w:footnoteReference w:id=...>``.

    Word footnotes are 1:1 with their reference, but Markdown lets one note be
    cited from several places (``attrs["index"]`` repeats). So, like pandoc, a
    fresh Word footnote is emitted per *reference* (content duplicated for
    repeats), in reference order, each with a unique id.
    """

    def __init__(self, document: DocxDocument, definitions: list[Node]) -> None:
        self._definitions = {
            int(item.get("attrs", {}).get("index", 0)): item for item in definitions
        }
        self._next_id = 1
        self._element = parse_xml(f'<w:footnotes xmlns:w="{_W_NS}"/>')
        self._element.append(_separator_footnote(-1, "separator"))
        self._element.append(_separator_footnote(0, "continuationSeparator"))
        package: OpcPackage = document.part.package
        self._part = XmlPart(
            PackURI(_FOOTNOTES_PARTNAME),
            _FOOTNOTES_CONTENT_TYPE,
            self._element,
            package,  # ty: ignore[invalid-argument-type]
        )
        document.part.relate_to(self._part, _FOOTNOTES_REL_TYPE)
        self._parent = _PartParent(self._part)

    def add_reference(self, paragraph: Paragraph, index: int) -> None:
        """Insert a superscript reference and emit its footnote definition."""
        item = self._definitions.get(index)
        if item is None:
            return
        footnote_id = self._next_id
        self._next_id += 1
        paragraph._p.append(_reference_run("w:footnoteReference", footnote_id))
        self._element.append(self._build_footnote(footnote_id, item))

    def _build_footnote(self, footnote_id: int, item: Node) -> Any:
        """Render a footnote definition into a ``<w:footnote>`` element.

        Footnotes are almost always a single paragraph; each paragraph-like
        block becomes a footnote paragraph, with the reference mark + a space
        leading the first one.
        """
        footnote = OxmlElement("w:footnote")
        footnote.set(qn("w:id"), str(footnote_id))
        for position, block in enumerate(item.get("children", [])):
            p_element = OxmlElement("w:p")
            p_pr = OxmlElement("w:pPr")
            p_style = OxmlElement("w:pStyle")
            p_style.set(qn("w:val"), _STYLE_ID_FOOTNOTE_TEXT)
            p_pr.append(p_style)
            p_element.append(p_pr)
            paragraph = Paragraph(p_element, self._parent)  # ty: ignore[invalid-argument-type]
            if position == 0:
                p_element.append(_reference_run("w:footnoteRef", None))
                paragraph.add_run(" ")
            _add_runs(paragraph, block.get("children", []), _Fmt(), None)
            footnote.append(p_element)
        return footnote


def _separator_footnote(footnote_id: int, separator_tag: str) -> Any:
    footnote = OxmlElement("w:footnote")
    footnote.set(qn("w:type"), separator_tag)
    footnote.set(qn("w:id"), str(footnote_id))
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    run.append(OxmlElement(f"w:{separator_tag}"))
    paragraph.append(run)
    footnote.append(paragraph)
    return footnote


def _reference_run(mark_tag: str, footnote_id: int | None) -> Any:
    """A run carrying the footnote-reference character style and the mark.

    ``mark_tag`` is ``w:footnoteReference`` (body, needs an id) or
    ``w:footnoteRef`` (the mark inside the note itself).
    """
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), _STYLE_ID_FOOTNOTE_REFERENCE)
    run_props.append(style)
    run.append(run_props)
    mark = OxmlElement(mark_tag)
    if footnote_id is not None:
        mark.set(qn("w:id"), str(footnote_id))
    run.append(mark)
    return run


# --------------------------------------------------------------------------- #
# Block-level rendering
# --------------------------------------------------------------------------- #
def _set_exact_line_spacing(paragraph_format: Any, line_pt: float) -> None:
    """Store a fixed line height. A float multiple uses Word auto metrics."""
    paragraph_format.line_spacing = Pt(line_pt)


def _drop_template_empty_paragraph(document: DocxDocument) -> None:
    """Remove the empty Normal paragraph python-docx ships in a new document."""
    body = document.element.body
    first = body.find(qn("w:p"))
    if first is None:
        return
    texts = first.findall(f".//{qn('w:t')}")
    if any((node.text or "").strip() for node in texts):
        return
    body.remove(first)


def _has_visible_inlines(children: list[Node] | None) -> bool:
    for child in children or []:
        kind = child.get("type")
        if kind in ("softbreak", "linebreak", "blank_line", "newline"):
            continue
        if kind == "text" and not str(child.get("raw") or "").strip():
            continue
        if kind == "text":
            return True
        if child.get("children") and _has_visible_inlines(child.get("children")):
            return True
        if kind not in ("text",):
            return True
    return False


def _set_paragraph_style(paragraph: Paragraph, style_name: str) -> None:
    """Set a paragraph's style by id, skipping python-docx's by-name lookup.

    ``add_paragraph(style=name)`` resolves the style by a linear, XML-parsing
    scan of every style and repeats it per paragraph, which dominates runtime on
    table/list-heavy documents. Built-in and added style ids are the spaceless
    form of the name, so set ``w:pStyle`` directly.
    """
    paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = style_name.replace(" ", "")


def _add_styled_paragraph(document: DocxDocument, style_name: str) -> Paragraph:
    paragraph = document.add_paragraph()
    _set_paragraph_style(paragraph, style_name)
    return paragraph


def _render_blocks(
    document: DocxDocument, nodes: list[Node], footnotes: "_Footnotes | None"
) -> None:
    # Like pandoc: the first prose paragraph after any non-paragraph block (a
    # heading, list, table, blockquote, code, or the document start) uses "First
    # Paragraph"; consecutive prose paragraphs use "Body Text".
    first_para_pending = True
    for node in nodes:
        node_type = node.get("type")
        if node_type in ("blank_line", "newline"):
            continue
        if node_type == "footnotes":
            # Definitions live in the footnotes part (emitted per reference), not
            # the body.
            continue
        if node_type in ("paragraph", "block_text"):
            children = node.get("children", [])
            if _is_image_only(children):
                _render_standalone_image(document, children)
                first_para_pending = False
            elif not _has_visible_inlines(children):
                continue
            else:
                style = _STYLE_FIRST_PARAGRAPH if first_para_pending else _STYLE_BODY
                paragraph = _add_styled_paragraph(document, style)
                _add_runs(paragraph, children, _Fmt(), footnotes)
                first_para_pending = False
            continue

        if node_type == "heading":
            level = min(int(node.get("attrs", {}).get("level", 1)), 6)
            paragraph = _add_styled_paragraph(document, f"Heading {level}")
            _add_runs(paragraph, node.get("children", []), _Fmt(), footnotes)
        elif node_type == "block_code":
            _render_code(document, node)
        elif node_type == "block_quote":
            _render_quote(document, node, footnotes)
        elif node_type == "list":
            _render_list(document, node, level=0, footnotes=footnotes)
        elif node_type == "thematic_break":
            _render_thematic_break(document)
        elif node_type == "table":
            _render_table(document, node, footnotes)
            # pandoc styles the paragraph after a table as Body Text, not First.
            first_para_pending = False
            continue
        elif "children" in node:
            # Unknown block wrapper: recurse so its content is not dropped.
            _render_blocks(document, node["children"], footnotes)
        first_para_pending = True


def _render_code(document: DocxDocument, node: Node) -> None:
    raw = str(node.get("raw", "")).rstrip("\n")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(4)
    _set_exact_line_spacing(paragraph.paragraph_format, WORD_CODE_LINE_PT)
    _shade_paragraph(paragraph, CODE_BG)
    for index, line in enumerate(raw.split("\n")):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = _MONOSPACE_FONT
        run.font.size = _CODE_FONT_SIZE


def _render_quote(
    document: DocxDocument, node: Node, footnotes: "_Footnotes | None"
) -> None:
    for child in node.get("children", []):
        if child.get("type") == "paragraph":
            paragraph = _add_styled_paragraph(document, _STYLE_BLOCK_TEXT)
            _add_paragraph_border(paragraph, "left", ACCENT)
            _add_runs(paragraph, child.get("children", []), _Fmt(), footnotes)
        else:
            _render_blocks(document, [child], footnotes)


def _is_image_only(children: list[Node]) -> bool:
    """True if the inline content is a single image (a standalone figure)."""
    meaningful = [
        child
        for child in children
        if child.get("type") not in ("softbreak", "linebreak")
        and not (child.get("type") == "text" and not str(child.get("raw", "")).strip())
    ]
    return len(meaningful) == 1 and meaningful[0].get("type") == "image"


def _render_standalone_image(document: DocxDocument, children: list[Node]) -> None:
    """Embed a standalone figure when bytes are present; else keep the caption."""
    image = next(child for child in children if child.get("type") == "image")
    alt = _collect_text(image.get("children", []))
    if image_bytes(image) is not None:
        paragraph = _add_styled_paragraph(document, _STYLE_BODY)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if _embed_picture(paragraph, image):
            if alt:
                _render_image_caption_text(document, alt)
            return
        _set_paragraph_style(paragraph, _STYLE_IMAGE_CAPTION)
        paragraph.add_run(alt or "image")
        return
    _render_image_caption_text(document, alt or "image")


def _render_image_caption_text(document: DocxDocument, alt: str) -> None:
    paragraph = _add_styled_paragraph(document, _STYLE_IMAGE_CAPTION)
    paragraph.add_run(alt)


def _embed_picture(paragraph: Paragraph, node: Node) -> bool:
    data = image_bytes(node)
    if data is None:
        return False
    size = fit_image_display_size(data)
    if size is None:
        return False
    width_in, height_in = size
    try:
        paragraph.add_run().add_picture(
            BytesIO(data), width=Inches(width_in), height=Inches(height_in)
        )
    except Exception:
        return False
    return True


def _render_list(
    document: DocxDocument,
    node: Node,
    level: int,
    footnotes: "_Footnotes | None",
) -> None:
    attrs = node.get("attrs", {})
    ordered = bool(attrs.get("ordered", False))
    # mistune exposes tight/loose as a top-level key on the list node.
    tight = bool(node.get("tight", True))
    style_level = min(level + 1, _MAX_LIST_LEVEL)
    list_style = "List Number" if ordered else "List Bullet"
    if style_level > 1:
        list_style = f"{list_style} {style_level}"
    continue_style = (
        "List Continue" if style_level == 1 else f"List Continue {style_level}"
    )
    start = int(attrs.get("start", 1))

    # Tight lists match pandoc's "Compact" style. Compact carries no list marker,
    # so numbering is applied directly from the built-in list style's definition
    # (which also supplies the indentation); each list gets its own instance so
    # ordered lists restart correctly. Loose lists keep the built-in list style.
    if tight:
        num_id = _create_list_numbering(document, list_style, start)
        item_style = _STYLE_COMPACT if num_id is not None else list_style
    else:
        # Ordered lists get their own numbering instance so each restarts at its
        # start value; bullets can rely on the built-in style.
        item_style = list_style
        num_id = (
            _create_list_numbering(document, list_style, start) if ordered else None
        )

    for item in node.get("children", []):
        if item.get("type") != "list_item":
            continue
        has_rendered_marker = False
        for child in item.get("children", []):
            child_type = child.get("type")
            if child_type in ("blank_line", "newline"):
                continue
            if child_type in ("block_text", "paragraph"):
                if not _has_visible_inlines(child.get("children")):
                    continue
                marker = not has_rendered_marker
                paragraph_style = item_style if marker else continue_style
                paragraph = _add_styled_paragraph(document, paragraph_style)
                # Indent 0.5" per level (matching pandoc); the marker line hangs.
                paragraph.paragraph_format.left_indent = Twips(
                    _LIST_INDENT_PER_LEVEL * (level + 1)
                )
                paragraph.paragraph_format.first_line_indent = Twips(
                    -_LIST_HANGING_INDENT if marker else 0
                )
                if marker and num_id is not None:
                    _apply_numbering(paragraph, num_id)
                _add_runs(paragraph, child.get("children", []), _Fmt(), footnotes)
                has_rendered_marker = True
            elif child_type == "list":
                _render_list(document, child, level + 1, footnotes)
            else:
                _render_blocks(document, [child], footnotes)


def _create_list_numbering(
    document: DocxDocument, list_style_name: str, start: int
) -> int | None:
    """Create a fresh numbering instance bound to a built-in list style.

    Returning a dedicated ``numId`` lets a paragraph in a non-list style (e.g.
    ``Compact``) still render list markers + indentation. A ``startOverride`` is
    always emitted (even for ``start == 1``): without it, multiple instances of
    the same abstract numbering continue one shared counter instead of each list
    restarting. Returns None if the style has no numbering definition (caller
    falls back to the list style).
    """
    style = document.styles[list_style_name]
    numbering = document.part.numbering_part.element
    abstract_num_id = _abstract_num_id_for_style(numbering, style.style_id)
    if abstract_num_id is None:
        return None

    num_ids = [
        int(num.get(qn("w:numId")))
        for num in numbering.findall(qn("w:num"))
        if num.get(qn("w:numId")) is not None
    ]
    next_num_id = max(num_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_num_id))

    abstract_num_id_el = OxmlElement("w:abstractNumId")
    abstract_num_id_el.set(qn("w:val"), abstract_num_id)
    num.append(abstract_num_id_el)

    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    lvl_override.append(start_override)
    num.append(lvl_override)

    numbering.append(num)
    return next_num_id


def _abstract_num_id_for_style(numbering: Any, style_id: str) -> str | None:
    for abstract_num in numbering.findall(qn("w:abstractNum")):
        for lvl in abstract_num.findall(qn("w:lvl")):
            p_style = lvl.find(qn("w:pStyle"))
            if p_style is not None and p_style.get(qn("w:val")) == style_id:
                return abstract_num.get(qn("w:abstractNumId"))
    return None


def _apply_numbering(paragraph: Paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = num_pr.get_or_add_ilvl()
    ilvl.val = 0
    num_id_el = num_pr.get_or_add_numId()
    num_id_el.val = num_id


def _render_thematic_break(document: DocxDocument) -> None:
    paragraph = document.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), RULE)
    borders.append(bottom)
    p_pr.append(borders)


def _render_table(
    document: DocxDocument, node: Node, footnotes: "_Footnotes | None"
) -> None:
    header_cells: list[Node] = []
    body_rows: list[list[Node]] = []
    for section in node.get("children", []):
        section_type = section.get("type")
        if section_type == "table_head":
            header_cells = section.get("children", [])
        elif section_type == "table_body":
            body_rows.extend(
                row.get("children", []) for row in section.get("children", [])
            )

    num_cols = len(header_cells) or (len(body_rows[0]) if body_rows else 0)
    if num_cols == 0:
        return

    table = document.add_table(rows=0, cols=num_cols)
    _set_table_full_width(table)
    _set_table_borders(table)
    _set_table_cell_margins(table)

    if header_cells:
        cells = table.add_row().cells
        _mark_header_row(table.rows[0])
        for index, cell_node in enumerate(header_cells[:num_cols]):
            _fill_cell(
                cells[index],
                cell_node,
                footnotes=footnotes,
                header=True,
            )
            _shade_cell(cells[index], TABLE_HEADER_FILL)
    for row_index, row in enumerate(body_rows):
        cells = table.add_row().cells
        _keep_row_together(table.rows[-1])
        for index, cell_node in enumerate(row[:num_cols]):
            _fill_cell(cells[index], cell_node, footnotes=footnotes, header=False)
            if row_index % 2 == 1:
                _shade_cell(cells[index], TABLE_ALT_FILL)

    _remove_fixed_cell_widths(table)


def _remove_fixed_cell_widths(table: Any) -> None:
    """Drop python-docx's fixed cell widths so columns auto-fit content.

    python-docx splits the full text width equally across columns; pandoc lets
    the table shrink to its content (leaving whitespace around small tables).
    """
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.tcPr
            if tc_pr is None:
                continue
            for tc_w in tc_pr.findall(qn("w:tcW")):
                tc_pr.remove(tc_w)


def _set_table_full_width(table: Any) -> None:
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    table._tbl.tblPr.append(width)


def _set_table_borders(table: Any) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:space"), "0")
        line.set(qn("w:color"), TABLE_BORDER)
        borders.append(line)
    table._tbl.tblPr.append(borders)


def _set_table_cell_margins(table: Any) -> None:
    """Compact cell padding: 40 twips vertical, 80 twips horizontal."""
    margins = OxmlElement("w:tblCellMar")
    for edge, width in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    table._tbl.tblPr.append(margins)


def _shade_cell(cell: _Cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _shade_paragraph(paragraph: Paragraph, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def _mark_header_row(row: Any) -> None:
    row_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    row_pr.append(header)
    _keep_row_together(row)


def _keep_row_together(row: Any) -> None:
    row_pr = row._tr.get_or_add_trPr()
    if row_pr.find(qn("w:cantSplit")) is None:
        row_pr.append(OxmlElement("w:cantSplit"))


_TABLE_CELL_ALIGN = {
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
}


def _fill_cell(
    cell: _Cell,
    cell_node: Node,
    footnotes: "_Footnotes | None",
    *,
    header: bool,
) -> None:
    paragraph = cell.paragraphs[0]
    _set_paragraph_style(paragraph, _STYLE_COMPACT)
    _set_exact_line_spacing(paragraph.paragraph_format, WORD_CELL_LINE_PT)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    alignment = _TABLE_CELL_ALIGN.get(str(cell_node.get("attrs", {}).get("align")))
    if alignment is not None:
        paragraph.alignment = alignment
    fmt = _Fmt(bold=header, color=_HEADER_COLOR if header else None)
    _add_runs(paragraph, cell_node.get("children", []), fmt, footnotes)
    for run in paragraph.runs:
        run.font.size = Pt(CELL_SIZE_PT)
        if header:
            run.bold = True
            run.font.color.rgb = _HEADER_COLOR


# --------------------------------------------------------------------------- #
# Inline rendering
# --------------------------------------------------------------------------- #
def _repair_paren_links(nodes: list[Node]) -> None:
    """Re-attach a balancing ``)`` that mistune split off a link destination.

    CommonMark allows balanced parentheses in a link/autolink destination (e.g.
    ``..._(novel)``), but mistune stops the URL at the first ``)`` and leaves it
    as following text. When a link's URL has unmatched ``(``, pull the matching
    ``)`` back from the next text node into the URL (and the visible text for an
    autolink), the way pandoc parses it.
    """
    for index, node in enumerate(nodes):
        if node.get("type") != "link" or index + 1 >= len(nodes):
            continue
        following = nodes[index + 1]
        if following.get("type") != "text":
            continue
        url = str(node.get("attrs", {}).get("url", ""))
        imbalance = url.count("(") - url.count(")")
        raw = str(following.get("raw", ""))
        take = 0
        while take < imbalance and take < len(raw) and raw[take] == ")":
            take += 1
        if take == 0:
            continue
        closing = ")" * take
        node.setdefault("attrs", {})["url"] = url + closing
        # For an autolink the visible text is the URL, and mistune split the same
        # ")" off it too. Detect that by the display ending with matching unclosed
        # "(" rather than comparing to the URL, which mistune may have re-encoded
        # (e.g. "\_" -> "%5C_"), so the closing reaches the display as well.
        children = node.get("children", [])
        if children and children[-1].get("type") == "text":
            display = str(children[-1].get("raw", ""))
            if display.count("(") - display.count(")") >= take:
                children[-1]["raw"] = display + closing
        following["raw"] = raw[take:]


def _add_runs(
    paragraph: Paragraph,
    nodes: list[Node],
    fmt: _Fmt,
    footnotes: "_Footnotes | None",
) -> None:
    _repair_paren_links(nodes)
    for node in nodes:
        node_type = node.get("type")
        if node_type == "text":
            # Decode HTML entities (e.g. ``&amp;``, ``&copy;``) the way a Markdown
            # renderer would; code spans below are intentionally left literal.
            _styled_run(paragraph, unescape(str(node.get("raw", ""))), fmt)
        elif node_type == "strong":
            _add_runs(
                paragraph, node.get("children", []), replace(fmt, bold=True), footnotes
            )
        elif node_type == "emphasis":
            _add_runs(
                paragraph,
                node.get("children", []),
                replace(fmt, italic=True),
                footnotes,
            )
        elif node_type == "strikethrough":
            _add_runs(
                paragraph,
                node.get("children", []),
                replace(fmt, strike=True),
                footnotes,
            )
        elif node_type == "codespan":
            _styled_run(paragraph, str(node.get("raw", "")), replace(fmt, code=True))
        elif node_type == "link":
            _add_hyperlink(paragraph, node, fmt, footnotes)
        elif node_type == "image":
            if not _embed_picture(paragraph, node):
                alt = _collect_text(node.get("children", []))
                _styled_run(paragraph, f"[image: {alt}]" if alt else "[image]", fmt)
        elif node_type == "footnote_ref":
            if footnotes is not None:
                footnotes.add_reference(
                    paragraph, int(node.get("attrs", {}).get("index", 0))
                )
        elif node_type == "softbreak":
            _styled_run(paragraph, " ", fmt)
        elif node_type == "linebreak":
            paragraph.add_run().add_break()
        elif node_type == "inline_html":
            _add_inline_html(paragraph, str(node.get("raw", "")))
        elif "children" in node:
            _add_runs(paragraph, node["children"], fmt, footnotes)
        elif "raw" in node:
            _styled_run(paragraph, unescape(str(node["raw"])), fmt)


def _styled_run(paragraph: Paragraph, text: str, fmt: _Fmt) -> None:
    run = paragraph.add_run(text)
    # Only set flags that are on, so plain runs don't emit redundant b="0"/i="0".
    if fmt.bold:
        run.bold = True
    if fmt.italic:
        run.italic = True
    if fmt.strike:
        run.font.strike = True
    if fmt.code:
        run.font.name = _MONOSPACE_FONT
    if fmt.color is not None:
        run.font.color.rgb = fmt.color


def _add_hyperlink(
    paragraph: Paragraph, node: Node, fmt: _Fmt, footnotes: "_Footnotes | None"
) -> None:
    url = str(node.get("attrs", {}).get("url", ""))
    children = node.get("children", [])
    if not url:
        _add_runs(paragraph, children, fmt, footnotes)
        return

    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    # Render the label's inline content so nested bold/italic/code survive, then
    # move those runs into the hyperlink and tag each with the "Hyperlink"
    # character style (colour, no underline) like pandoc.
    start = len(paragraph._p)
    _add_runs(paragraph, children, fmt, footnotes)
    rendered = paragraph._p[start:]
    for element in rendered:
        paragraph._p.remove(element)
        if element.tag == qn("w:r"):
            _apply_hyperlink_style(element)
        hyperlink.append(element)
    if not rendered:
        # Empty label: fall back to the URL as the visible text.
        hyperlink.append(_hyperlink_text_run(url))
    paragraph._p.append(hyperlink)


def _apply_hyperlink_style(run: Any) -> None:
    run_props = run.find(qn("w:rPr"))
    if run_props is None:
        run_props = OxmlElement("w:rPr")
        run.insert(0, run_props)
    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), _STYLE_ID_HYPERLINK)
    run_props.insert(0, run_style)  # rStyle is first in CT_RPr


def _hyperlink_text_run(text: str) -> Any:
    run = OxmlElement("w:r")
    _apply_hyperlink_style(run)
    text_el = OxmlElement("w:t")
    text_el.text = text
    if text != text.strip():
        text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)
    return run


def _add_inline_html(paragraph: Paragraph, raw: str) -> None:
    normalized = raw.strip().lower()
    if normalized in ("<br>", "<br/>", "<br />"):
        paragraph.add_run().add_break()


def _collect_text(nodes: list[Node]) -> str:
    """Flatten inline nodes to plain text (for link labels and image alt text)."""
    parts: list[str] = []
    for node in nodes:
        if node.get("type") == "text":
            parts.append(unescape(str(node.get("raw", ""))))
        elif node.get("children"):
            parts.append(_collect_text(node["children"]))
        elif "raw" in node:
            parts.append(unescape(str(node["raw"])))
    return "".join(parts)
