"""Build catalog Word report templates from a per-slug spec.

``generate_official_docx(slug)`` returns the document bytes that sync
attaches. The file is a layout reference: A4 research-note styles plus
the placeholder tokens the agent must fill.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Final

from docx import Document

from onyx.system_catalog.builtin.word import styles


@dataclass(frozen=True)
class TableColumn:
    header: str
    field: str
    description: str
    example: str = ""


@dataclass(frozen=True)
class ScalarField:
    name: str
    description: str
    example: str = ""
    heading: str | None = None


@dataclass(frozen=True)
class SectionBlock:
    """One report section: heading plus an optional body token and/or table."""

    heading: str
    body_token: str | None = None
    table_collection: str = ""
    table_columns: tuple[TableColumn, ...] = ()


@dataclass(frozen=True)
class OfficialWordSpec:
    slug: str
    title: str
    # Legacy single-table layout fields. Specs using ``sections`` leave these
    # empty; the sections carry their own tables inline.
    table_heading: str = ""
    table_collection: str = ""
    table_columns: tuple[TableColumn, ...] = ()
    extra_scalars: tuple[ScalarField, ...] = ()
    extra_sections: tuple[tuple[str, str], ...] = ()
    # Long-form specs declare sections inline (text and tables interleaved in
    # order) and skip the trailing single-table layout.
    sections: tuple[SectionBlock, ...] = ()


_FINDINGS: Final[tuple[TableColumn, ...]] = (
    TableColumn("发现", "title", "一条发现的标题", "进项认证异常"),
    TableColumn("严重程度", "severity", "高 / 中 / 低", "高"),
    TableColumn("金额", "amount", "涉及金额，无金额写 未获取", "120万"),
    TableColumn("责任人", "owner", "跟进责任人", "税务经理"),
    TableColumn("来源", "source", "证据来源或文件名", "增值税申报表"),
)

_EXCEPTIONS: Final[tuple[TableColumn, ...]] = (
    TableColumn("例外", "title", "例外事项标题", "银行未达账"),
    TableColumn("严重程度", "severity", "高 / 中 / 低", "中"),
    TableColumn("金额", "amount", "涉及金额", "35万"),
    TableColumn("责任人", "owner", "跟进责任人", "出纳"),
    TableColumn("来源", "source", "对账文件或模块", "工商银行对账单"),
)

_DIFFERENCES: Final[tuple[TableColumn, ...]] = (
    TableColumn("差异", "title", "差异事项", "发票未入账"),
    TableColumn("严重程度", "severity", "高 / 中 / 低", "中"),
    TableColumn("金额", "amount", "差异金额", "8.2万"),
    TableColumn("责任人", "owner", "跟进责任人", "应付会计"),
    TableColumn("来源", "source", "两侧文件", "供应商对账单"),
)

_ACTIONS: Final[tuple[TableColumn, ...]] = (
    TableColumn("行动项", "title", "具体动作", "补充关账证据"),
    TableColumn("负责人", "owner", "行动项负责人", "财务经理"),
    TableColumn("截止日期", "due", "YYYY-MM-DD，未指定写 未指定", "2026-09-15"),
    TableColumn("状态", "status", "未开始 / 进行中 / 完成", "未开始"),
    TableColumn("备注", "source", "依据或会议决议", "第 3 项决议"),
)

_SLIDES: Final[tuple[TableColumn, ...]] = (
    TableColumn("页", "title", "幻灯片标题", "问题与结论"),
    TableColumn("要点", "severity", "该页要说清的一点", "关账仍被银行未达阻塞"),
    TableColumn("证据", "amount", "支撑数字或来源", "未达 35 万"),
    TableColumn("受众", "owner", "该页主要给谁看", "管理层"),
    TableColumn("备注", "source", "演讲提示", "先结论后细节"),
)


def _spec(
    slug: str,
    title: str,
    *,
    table_heading: str = "发现与例外",
    table_collection: str = "findings",
    table_columns: tuple[TableColumn, ...] = _FINDINGS,
    extra_scalars: tuple[ScalarField, ...] = (),
    extra_sections: tuple[tuple[str, str], ...] = (),
) -> OfficialWordSpec:
    return OfficialWordSpec(
        slug=slug,
        title=title,
        table_heading=table_heading,
        table_collection=table_collection,
        table_columns=table_columns,
        extra_scalars=extra_scalars,
        extra_sections=extra_sections,
    )


_KPI: Final[tuple[TableColumn, ...]] = (
    TableColumn("指标", "metric", "指标名称与单位", "毛利率（%）"),
    TableColumn("年度1", "y1", "第一个分析年度数值", "46.09"),
    TableColumn("年度2", "y2", "第二个分析年度数值", "34.54"),
    TableColumn("年度3", "y3", "第三个分析年度数值", "37.96"),
    TableColumn("年度4", "y4", "第四个分析年度数值", "32.17"),
    TableColumn("年度5", "y5", "第五个分析年度数值", "35.60"),
)

_RISK_MATRIX: Final[tuple[TableColumn, ...]] = (
    TableColumn("编号", "risk_id", "TX-01～TX-10 或 OP-01～OP-08", "TX-01"),
    TableColumn("风险点", "risk", "风险名称", "实际所得税率突破法定税率"),
    TableColumn("等级", "level", "高 / 中 / 低", "高"),
    TableColumn("风险评分", "score", "高=56 中=20 低=4", "56"),
    TableColumn("处置建议", "action", "立即治理 / 常规跟进", "立即治理"),
)

_RECOMMENDATIONS: Final[tuple[TableColumn, ...]] = (
    TableColumn("建议项", "item", "建议名称", "税收优惠资格续期专项"),
    TableColumn("优先级", "priority", "高 / 中 / 低", "高"),
    TableColumn("实施周期", "timeline", "预计完成时间", "90日内"),
    TableColumn(
        "落地措施与预期收益",
        "measures",
        "具体动作与量化收益",
        "复审前完成研发归集整改，预计年节税约1300万",
    ),
)

_PROCEDURES: Final[tuple[TableColumn, ...]] = (
    TableColumn("风险领域", "area", "对应风险条目", "收入确认"),
    TableColumn("等级", "level", "高 / 中 / 低", "中高"),
    TableColumn("核心问题", "issue", "要核实的问题", "期末集中确认的合同资产可回收性"),
    TableColumn(
        "建议程序", "procedure", "审计导向的核查动作", "函证、期后回款、截止测试"
    ),
)

_SUBSEQUENT: Final[tuple[TableColumn, ...]] = (
    TableColumn("事项", "event", "期后事件", "定向增发终止"),
    TableColumn("公开信息", "disclosure", "公告或新闻来源", "巨潮公告 2026-02-11"),
    TableColumn("分析影响", "impact", "对风险判断的影响", "补流计划落空，关注流动性"),
)

_FINANCE_TAX_RISK_REPORT: Final[OfficialWordSpec] = OfficialWordSpec(
    slug="finance_tax_risk_report",
    title="财税与经营风险分析报告",
    sections=(
        SectionBlock(heading="一、报告说明与数据来源", body_token="{{report_notes}}"),
        SectionBlock(
            heading="二、老板数据看板（核心 KPI 速览）",
            body_token="{{kpi_insights}}",
            table_collection="kpi",
            table_columns=_KPI,
        ),
        SectionBlock(heading="三、经营与财务趋势分析", body_token="{{trend_analysis}}"),
        SectionBlock(
            heading="四、五年财务数据透视", body_token="{{statements_analysis}}"
        ),
        SectionBlock(
            heading="五、税务与现金专项分析", body_token="{{tax_cash_analysis}}"
        ),
        SectionBlock(heading="六、财税风险分析及识别点", body_token="{{tax_risks}}"),
        SectionBlock(heading="七、经营风险识别", body_token="{{operating_risks}}"),
        SectionBlock(heading="八、重大专题与治理风险", body_token="{{special_topics}}"),
        SectionBlock(
            heading="九、专业改进建议",
            body_token="{{recommendations}}",
            table_collection="recommendations",
            table_columns=_RECOMMENDATIONS,
        ),
        SectionBlock(
            heading="十、风险矩阵与综合评级",
            body_token="{{overall_rating}}",
            table_collection="risk_matrix",
            table_columns=_RISK_MATRIX,
        ),
        SectionBlock(heading="十一、优先级整改建议", body_token="{{remediation_plan}}"),
        SectionBlock(
            heading="风险清单及建议核查程序",
            table_collection="procedures",
            table_columns=_PROCEDURES,
        ),
        SectionBlock(
            heading="期后事项与外部信息",
            body_token="{{subsequent_events}}",
            table_collection="subsequent_events",
            table_columns=_SUBSEQUENT,
        ),
        SectionBlock(heading="十二、结论", body_token="{{conclusion}}"),
        SectionBlock(heading="未决问题", body_token="{{open_issues}}"),
        SectionBlock(heading="数据缺口", body_token="{{data_gaps}}"),
        SectionBlock(
            heading="附录与免责声明", body_token="{{appendix_and_disclaimer}}"
        ),
    ),
)

OFFICIAL_WORD_SPECS: Final[dict[str, OfficialWordSpec]] = {
    spec.slug: spec
    for spec in (
        _spec("compliance_risk", "合规风险预警报告"),
        _spec("policy_trend", "税收政策趋势研判"),
        _spec(
            "monthly_close",
            "月度关账报告",
            table_heading="例外事项",
            table_collection="exceptions",
            table_columns=_EXCEPTIONS,
        ),
        _spec("financial_review", "财务报表分析报告"),
        _spec("target_landscape", "靶点文献与竞争格局"),
        _spec("patent_fto", "专利自由实施评估"),
        _spec("clinical_pipeline", "临床管线扫描"),
        _spec("cmc_quality", "CMC 与质量风险简报"),
        _spec("regulatory_pathway", "注册申报路径"),
        _spec("research_brief", "调研简报"),
        _spec("data_analysis_report", "数据分析报告"),
        _spec(
            "meeting_minutes",
            "会议纪要",
            table_heading="行动项",
            table_collection="actions",
            table_columns=_ACTIONS,
            extra_scalars=(
                ScalarField(
                    "meeting_title",
                    "会议主题",
                    "9 月关账评审",
                    "会议主题",
                ),
                ScalarField(
                    "attendees",
                    "参会人，逗号分隔",
                    "张三, 李四",
                    "参会人",
                ),
            ),
        ),
        _spec("initiation_report", "药物立项评估报告"),
        _spec("listed_company_audit", "上市公司财务审计报告"),
        _spec(
            "supplier_recon",
            "供应商对账报告",
            table_heading="差异清单",
            table_collection="differences",
            table_columns=_DIFFERENCES,
        ),
        _spec(
            "bank_ledger",
            "银行流水核对报告",
            table_heading="未达账项",
            table_collection="differences",
            table_columns=_DIFFERENCES,
        ),
        _spec("expense_rollup", "费用归集报告"),
        _spec(
            "invoice_compliance",
            "发票合规审查报告",
            table_heading="例外清单",
            table_collection="exceptions",
            table_columns=_EXCEPTIONS,
        ),
        _spec("opex_variance", "费用波动分析报告"),
        _spec("ar_risk", "应收账款风险清单"),
        _spec(
            "deck_outline",
            "演示文稿大纲",
            table_heading="幻灯片大纲",
            table_collection="slides",
            table_columns=_SLIDES,
            extra_scalars=(
                ScalarField(
                    "audience",
                    "主要受众",
                    "管理层",
                    "受众",
                ),
                ScalarField(
                    "key_message",
                    "一句话结论",
                    "本月不可关账",
                    "关键信息",
                ),
            ),
        ),
        _FINANCE_TAX_RISK_REPORT,
    )
}


def build_document(spec: OfficialWordSpec) -> bytes:
    document = Document()
    styles.configure_document(document)
    styles.add_header(document)
    styles.add_footer(document)
    styles.add_title(document, spec.title)
    styles.add_cover_line(document)
    for extra in spec.extra_scalars:
        if extra.heading:
            styles.add_heading(document, extra.heading)
            styles.add_body(document, "{{" + extra.name + "}}")
    if spec.sections:
        for section in spec.sections:
            styles.add_heading(document, section.heading)
            if section.body_token:
                styles.add_body(document, section.body_token)
            if section.table_columns:
                styles.add_bordered_table(
                    document,
                    headers=[column.header for column in section.table_columns],
                    prototype_cells=[
                        "{{" + section.table_collection + "." + column.field + "}}"
                        for column in section.table_columns
                    ],
                )
    else:
        for heading, token in spec.extra_sections:
            styles.add_heading(document, heading)
            styles.add_body(document, token)
        styles.add_heading(document, spec.table_heading)
        styles.add_bordered_table(
            document,
            headers=[column.header for column in spec.table_columns],
            prototype_cells=[
                "{{" + spec.table_collection + "." + column.field + "}}"
                for column in spec.table_columns
            ],
        )
        styles.add_closing(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_official_docx(slug: str) -> bytes:
    """Return generated Word bytes for ``slug``."""
    spec = OFFICIAL_WORD_SPECS.get(slug)
    if spec is None:
        raise KeyError(f"No official Word builder for '{slug}'")
    return build_document(spec)


def official_builder_slugs() -> frozenset[str]:
    return frozenset(OFFICIAL_WORD_SPECS)
