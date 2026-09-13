"""Professional Word layout for catalog report templates.

This is the visual contract for generated ``.docx`` assets. It follows
compact A4 research-note practice (sans body, ink hierarchy, charcoal
tables) and shares tokens with Markdown export.
"""

from __future__ import annotations

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from onyx.server.features.build.session.md_export_style import (
    ACCENT,
    BODY_EAST_ASIA,
    BODY_FONT,
    BODY_LINE_SPACING,
    BODY_SIZE_PT,
    BODY_SPACE_AFTER_PT,
    CELL_SIZE_PT,
    FOOTER_DISTANCE_MM,
    HEADER_DISTANCE_MM,
    HEADING_EAST_ASIA,
    HEADING_FONT,
    INK,
    MARGIN_BOTTOM_MM,
    MARGIN_LEFT_MM,
    MARGIN_RIGHT_MM,
    MARGIN_TOP_MM,
    MUTED,
    PAGE_HEIGHT_MM,
    PAGE_WIDTH_MM,
    RULE,
    TABLE_ALT_FILL,
    TABLE_BORDER,
    TABLE_HEADER_FILL,
    WHITE,
)


def _rgb(hex6: str) -> RGBColor:
    return RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))


def set_run_font(
    run: Run,
    *,
    size_pt: float = BODY_SIZE_PT,
    bold: bool = False,
    color: RGBColor | None = None,
    heading: bool = False,
) -> None:
    latin = HEADING_FONT if heading else BODY_FONT
    east_asia = HEADING_EAST_ASIA if heading else BODY_EAST_ASIA
    run.font.name = latin
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def configure_document(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width = Mm(PAGE_WIDTH_MM)
    section.page_height = Mm(PAGE_HEIGHT_MM)
    section.top_margin = Mm(MARGIN_TOP_MM)
    section.bottom_margin = Mm(MARGIN_BOTTOM_MM)
    section.left_margin = Mm(MARGIN_LEFT_MM)
    section.right_margin = Mm(MARGIN_RIGHT_MM)
    section.header_distance = Mm(HEADER_DISTANCE_MM)
    section.footer_distance = Mm(FOOTER_DISTANCE_MM)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE_PT)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.line_spacing = BODY_LINE_SPACING
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BODY_SPACE_AFTER_PT)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_fonts.set(qn("w:eastAsia"), BODY_EAST_ASIA)


def add_header(document: DocumentObject) -> None:
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run("{{company_header}}")
    set_run_font(run, size_pt=8, color=_rgb(MUTED))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_edge_border(paragraph, "bottom", RULE)


def add_footer(document: DocumentObject) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_edge_border(paragraph, "top", RULE)
    prefix = paragraph.add_run("— ")
    set_run_font(prefix, size_pt=8, color=_rgb(MUTED))
    _add_page_field(paragraph)
    suffix = paragraph.add_run(" —")
    set_run_font(suffix, size_pt=8, color=_rgb(MUTED))


def add_title(document: DocumentObject, title: str) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run(title)
    set_run_font(run, size_pt=18, bold=True, color=_rgb(INK), heading=True)
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True
    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(10)
    _add_edge_border(rule, "bottom", ACCENT, size="18", space="1")


def add_cover_line(document: DocumentObject) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(
        "主体  {{entity_name}}    期间  {{period}}    "
        "币种  {{currency}}    编制  {{report_date}}"
    )
    set_run_font(run, size_pt=9, color=_rgb(MUTED))
    paragraph.paragraph_format.space_after = Pt(14)


def add_heading(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size_pt=13, bold=True, color=_rgb(ACCENT), heading=True)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True


def add_body(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size_pt=BODY_SIZE_PT, color=_rgb(INK))


def add_bordered_table(
    document: DocumentObject,
    headers: list[str],
    prototype_cells: list[str],
) -> None:
    table = document.add_table(rows=2, cols=len(headers))
    _set_table_full_width(table)
    _set_table_borders(table)
    _set_table_cell_margins(table)
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(header)
        set_run_font(
            run, size_pt=CELL_SIZE_PT, bold=True, color=_rgb(WHITE), heading=True
        )
        _shade_cell(cell, TABLE_HEADER_FILL)
    for index, token in enumerate(prototype_cells):
        cell = table.cell(1, index)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(token)
        set_run_font(run, size_pt=CELL_SIZE_PT, color=_rgb(INK))
        _shade_cell(cell, TABLE_ALT_FILL)


def add_closing(document: DocumentObject) -> None:
    add_heading(document, "结论")
    add_body(document, "{{conclusion}}")
    add_heading(document, "未决问题")
    add_body(document, "{{open_issues}}")
    add_heading(document, "数据缺口")
    add_body(document, "{{data_gaps}}")


def _add_edge_border(
    paragraph: Paragraph, edge: str, color: str, size: str = "6", space: str = "4"
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    line = OxmlElement(f"w:{edge}")
    line.set(qn("w:val"), "single")
    line.set(qn("w:sz"), size)
    line.set(qn("w:space"), space)
    line.set(qn("w:color"), color)
    borders.append(line)


def _set_table_full_width(table: object) -> None:
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    table._tbl.tblPr.append(width)  # type: ignore[attr-defined]


def _set_table_borders(table: object) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:space"), "0")
        line.set(qn("w:color"), TABLE_BORDER)
        borders.append(line)
    table._tbl.tblPr.append(borders)  # type: ignore[attr-defined]


def _set_table_cell_margins(table: object) -> None:
    margins = OxmlElement("w:tblCellMar")
    for edge, width in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    table._tbl.tblPr.append(margins)  # type: ignore[attr-defined]


def _shade_cell(cell: object, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)  # type: ignore[attr-defined]


def _add_page_field(paragraph: Paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, size_pt=8, color=_rgb(MUTED))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
