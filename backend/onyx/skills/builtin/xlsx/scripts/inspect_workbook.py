"""Print a JSON summary of an XLSX workbook.

Example:
    python .opencode/skills/xlsx/scripts/inspect_workbook.py --input workbook.xlsx --max-rows 5
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Type detection reads past the display sample so a column whose first rows are
# blank is still classified, without paging through a huge sheet.
TYPE_SCAN_LIMIT = 1_000


def json_value(value: Any) -> Any:
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return value


def type_name(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int | float):
        return "number"
    if isinstance(value, datetime | date | time):
        return "date"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def summarize_sheet(sheet: Any, max_rows: int) -> dict[str, Any]:
    rows = [
        list(row) for row in sheet.iter_rows(max_row=TYPE_SCAN_LIMIT, values_only=True)
    ]
    headers = [json_value(value) for value in rows[0]] if rows else []
    sample_rows = [
        [json_value(value) for value in row] for row in rows[1 : max_rows + 1]
    ]

    column_types: dict[str, list[str]] = {}
    for column_index, header in enumerate(headers):
        counts = Counter(
            type_name(row[column_index]) for row in rows[1:] if column_index < len(row)
        )
        name = (
            str(header)
            if header not in (None, "")
            else get_column_letter(column_index + 1)
        )
        column_types[name] = [item for item, _ in counts.most_common()]

    return {
        "name": sheet.title,
        "dimensions": {
            "rows": sheet.max_row or 0,
            "columns": sheet.max_column or 0,
        },
        "header_row": headers,
        "column_types": column_types,
        "sample_rows": sample_rows,
        "scanned_rows": len(rows),
    }


def inspect_workbook(path: Path, max_rows: int) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"input not found: {path}")
    if max_rows < 0:
        raise ValueError("--max-rows must be zero or greater")

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return {
            "input": str(path),
            "sheet_names": workbook.sheetnames,
            "sheets": [
                summarize_sheet(workbook[name], max_rows)
                for name in workbook.sheetnames
            ],
        }
    finally:
        workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inspect an XLSX workbook and print JSON"
    )
    parser.add_argument("--input", required=True, help="Input workbook")
    parser.add_argument("--max-rows", type=int, default=5, help="Sample rows per sheet")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = inspect_workbook(Path(args.input), args.max_rows)
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        sys.exit(1)
