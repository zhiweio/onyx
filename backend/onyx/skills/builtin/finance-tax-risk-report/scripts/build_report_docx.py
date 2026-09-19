"""Render the finance-tax-risk report markdown into a styled .docx.

Markdown subset (see ``assets/report_template.md``):

- YAML frontmatter: title / company / ticker / period / data_scope /
  analysis_date / disclaimer — rendered as the cover page.
- ``#`` chapters, ``##`` sections, ``###`` subsections. Numbers are stripped
  and re-applied in sequence (一、二、… / N.M / （一）…), which fixes the
  duplicated and skipped chapter numbers of hand-assembled reports.
- GFM tables, ``-`` bullets, numbered lines (kept as literal text), paragraphs
  with ``**bold**``.
- ``> [!风险] 标题`` / ``> [!洞察] 标题`` callouts become shaded paragraphs,
  not 1x1 tables.
- ``![图题](path)`` embeds the image; missing files render as a placeholder.
- A real Word TOC field follows the cover, so headings are navigable.

Usage::

    python build_report_docx.py --input report.md --output report.docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x2E, 0x5E, 0x8C)
RISK_FILL = "FDECEA"
RISK_BORDER = "C0392B"
INSIGHT_FILL = "EEF2F7"
INSIGHT_BORDER = "2E5E8C"
HEADER_FILL = "E8EDF3"
CN_NUMERALS = "一二三四五六七八九十"

_HEADING_STRIP = (
    re.compile(r"^[一二三四五六七八九十百]+、\s*"),
    re.compile(r"^\d+(?:\.\d+)*\s*"),
    re.compile(r"^（[一二三四五六七八九十]+）\s*"),
    re.compile(r"^\d+[．.、]\s*"),
)
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
_CALLOUT = re.compile(r"^>\s*\[!(风险|洞察)\]\s*(.*)$")
_FRONTMATTER = re.compile(r"^---\s*$")


def _cn_numeral(index: int) -> str:
    if index <= 10:
        return CN_NUMERALS[index - 1]
    if index < 20:
        return f"十{CN_NUMERALS[index - 11]}"
    tens, ones = divmod(index, 10)
    if ones == 0:
        return f"{CN_NUMERALS[tens - 1]}十"
    return f"{CN_NUMERALS[tens - 1]}十{CN_NUMERALS[ones - 1]}"


def _strip_numbering(text: str) -> str:
    for pattern in _HEADING_STRIP:
        text = pattern.sub("", text, count=1)
    return text.strip()


def _set_east_asia(run: object, latin: str, east_asia: str) -> None:
    r_pr = run._element.get_or_add_rPr()  # type: ignore[attr-defined]
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _set_style_east_asia(style_element: object, latin: str, east_asia: str) -> None:
    r_pr = style_element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _configure_styles(document: DocumentObject) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    _set_style_east_asia(normal.element, "Calibri", "宋体")
    specs = (("Heading 1", 16, INK), ("Heading 2", 14, ACCENT), ("Heading 3", 12, INK))
    for name, size, color in specs:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        _set_style_east_asia(style.element, "Calibri", "黑体")


def _add_runs_with_bold(paragraph: object, text: str) -> None:
    position = 0
    for match in _BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])  # type: ignore[attr-defined]
        bold = paragraph.add_run(match.group(1))  # type: ignore[attr-defined]
        bold.bold = True  # type: ignore[attr-defined]
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])  # type: ignore[attr-defined]


def _add_page_break(document: DocumentObject) -> None:
    from docx.enum.text import WD_BREAK

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)  # type: ignore[attr-defined]


def _add_toc_field(document: DocumentObject) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "目录将在打开文档后生成（Word 中按 F9 更新域）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, hint, end):
        run._r.append(element)  # type: ignore[attr-defined]


def _add_footer_page_number(document: DocumentObject) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, end):
        run._r.append(element)  # type: ignore[attr-defined]


def _shade_paragraph(paragraph: object, fill: str, border_color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), border_color)
    borders.append(left)
    p_pr.append(borders)


def _add_table(document: DocumentObject, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    table = document.add_table(rows=len(padded), cols=width)
    table.style = "Table Grid"
    for r_index, row in enumerate(padded):
        for c_index, cell_text in enumerate(row):
            cell = table.cell(r_index, c_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _add_runs_with_bold(paragraph, cell_text)
            if r_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = INK
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), HEADER_FILL)
                cell._tc.get_or_add_tcPr().append(shading)


def _add_image(
    document: DocumentObject, caption: str, path: str, base_dir: Path
) -> None:
    candidates = [Path(path), base_dir / path]
    resolved = next((c for c in candidates if c.is_file()), None)
    if resolved is None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"（图未生成：{path}）")
        run.italic = True
        run.font.color.rgb = MUTED
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(resolved), width=Mm(150))  # type: ignore[attr-defined]
    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption_paragraph.add_run(caption)
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def _parse_frontmatter(lines: list[str]) -> tuple[dict[str, str], int]:
    if not lines or not _FRONTMATTER.match(lines[0]):
        return {}, 0
    metadata: dict[str, str] = {}
    for index in range(1, len(lines)):
        if _FRONTMATTER.match(lines[index]):
            return metadata, index + 1
        line = lines[index]
        if ":" in line and not line.startswith((" ", "\t")):
            key, _, value = line.partition(":")
            metadata[key.strip()] = value.strip().strip('"')
    return metadata, len(lines)


def _render_cover_and_toc(document: DocumentObject, metadata: dict[str, str]) -> None:
    title = metadata.get("title", "财税与经营风险分析报告")
    company = metadata.get("company", "")
    cover_paragraph = document.add_paragraph()
    cover_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_paragraph.paragraph_format.space_before = Pt(120)
    run = cover_paragraph.add_run(company + title if company else title)
    run.font.size = Pt(24)
    run.bold = True
    run.font.color.rgb = INK
    _set_east_asia(run, "Calibri", "黑体")
    for key in ("ticker", "period", "data_scope", "analysis_date"):
        if metadata.get(key):
            line = document.add_paragraph()
            line.alignment = WD_ALIGN_PARAGRAPH.CENTER
            line_run = line.add_run(metadata[key])
            line_run.font.size = Pt(11)
            line_run.font.color.rgb = MUTED
    if metadata.get("disclaimer"):
        disclaimer = document.add_paragraph()
        disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        disclaimer_run = disclaimer.add_run(metadata["disclaimer"])
        disclaimer_run.font.size = Pt(9)
        disclaimer_run.font.color.rgb = MUTED
    _add_page_break(document)

    toc_title = document.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_run = toc_title.add_run("目  录")
    toc_run.bold = True
    toc_run.font.size = Pt(16)
    _add_toc_field(document)
    _add_page_break(document)


def _collect_table(
    lines: list[str], start: int, first: str
) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    stripped = first
    while True:
        cells = [c.strip() for c in stripped.strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
            rows.append(cells)
        if index >= len(lines) or not lines[index].strip().startswith("|"):
            return rows, index
        stripped = lines[index].strip()
        index += 1


def render(markdown_path: Path, output_path: Path) -> dict[str, int]:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    metadata, body_start = _parse_frontmatter(lines)
    base_dir = markdown_path.parent

    document = Document()
    _configure_styles(document)
    _add_footer_page_number(document)
    _render_cover_and_toc(document, metadata)

    stats = {"chapters": 0, "tables": 0, "images": 0, "callouts": 0}
    chapter = 0
    section = 0
    subsection = 0

    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            _add_runs_with_bold(paragraph, " ".join(paragraph_buffer).strip())
            paragraph_buffer = []

    index = body_start
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        index += 1

        if not stripped or stripped.startswith("<!--"):
            flush_paragraph()
            continue
        if stripped == "-->":
            continue
        if _FRONTMATTER.match(stripped):
            continue
        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            text = _strip_numbering(stripped.lstrip("#"))
            if level == 1:
                chapter += 1
                section = 0
                text = f"{_cn_numeral(chapter)}、{text}"
                stats["chapters"] += 1
            elif level == 2:
                section += 1
                subsection = 0
                text = f"{chapter}.{section} {text}"
            else:
                subsection += 1
                text = f"（{_cn_numeral(subsection)}）{text}"
            document.add_heading(text, level=min(level, 3))
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            table_rows, index = _collect_table(lines, index, stripped)
            _add_table(document, table_rows)
            stats["tables"] += 1
            continue
        callout = _CALLOUT.match(stripped)
        if callout:
            flush_paragraph()
            kind, heading = callout.group(1), callout.group(2)
            body: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                body.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = document.add_paragraph()
            _shade_paragraph(
                paragraph,
                RISK_FILL if kind == "风险" else INSIGHT_FILL,
                RISK_BORDER if kind == "风险" else INSIGHT_BORDER,
            )
            title_run = paragraph.add_run(f"【{kind}】{heading}")
            title_run.bold = True
            if body:
                paragraph.add_run("\n" + " ".join(piece for piece in body if piece))
            stats["callouts"] += 1
            continue
        image = _IMAGE.match(stripped)
        if image:
            flush_paragraph()
            _add_image(document, image.group(1), image.group(2), base_dir)
            stats["images"] += 1
            continue
        if stripped.startswith(("- ", "* ")):
            flush_paragraph()
            bullet = document.add_paragraph(style="List Bullet")
            _add_runs_with_bold(bullet, stripped[2:])
            continue
        if re.match(r"^\d+[．.、)]\s*", stripped):
            flush_paragraph()
            numbered = document.add_paragraph()
            numbered.paragraph_format.left_indent = Mm(6)
            _add_runs_with_bold(numbered, stripped)
            continue
        paragraph_buffer.append(stripped)
    flush_paragraph()

    document.save(str(output_path))
    return stats


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", type=Path, required=True, help="报告 markdown 路径")
    parser.add_argument("--output", type=Path, required=True, help="输出 docx 路径")
    args = parser.parse_args()
    stats = render(args.input, args.output)
    print(
        f"written: {args.output} "
        f"(chapters={stats['chapters']}, tables={stats['tables']}, "
        f"images={stats['images']}, callouts={stats['callouts']})"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
