#!/usr/bin/env python3
"""Deterministic document ingest for Craft long jobs.

Standalone so the sandbox can run it without the onyx package.
Writes MANIFEST.json, extracted JSON, and an exception CSV.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import subprocess
import sys
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

TABLE_EXTS = {".xlsx", ".xls", ".csv"}
DOCX_EXTS = {".docx"}
PPTX_EXTS = {".pptx", ".ppt"}
PDF_EXTS = {".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}
SKIP_DIR_NAMES = {".git", ".opencode", "node_modules", "__pycache__"}
TEXT_PREVIEW_CHARS = 4000
PDF_TEXT_MIN_CHARS = 40
VISION_BATCH_PAGES = 4


@dataclass
class ManifestEntry:
    file_id: str
    path: str
    parser: str
    status: str
    confidence: float
    error: str | None = None
    extracted_path: str | None = None
    page_count: int | None = None
    sheet_count: int | None = None
    needs_vision: bool = False


@dataclass
class IngestResult:
    entries: list[ManifestEntry] = field(default_factory=list)

    def to_manifest(self) -> dict[str, Any]:
        return {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "files": [asdict(entry) for entry in self.entries],
        }


def classify_parser(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in TABLE_EXTS:
        return "pandas"
    if ext in DOCX_EXTS:
        return "python-docx"
    if ext in PPTX_EXTS:
        return "python-pptx"
    if ext in PDF_EXTS:
        return "pdfplumber"
    if ext in IMAGE_EXTS:
        return "vision"
    return "skip"


def file_id_for(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(str(path).encode("utf-8", errors="replace"))
    try:
        digest.update(str(path.stat().st_size).encode())
    except OSError:
        pass
    return digest.hexdigest()[:16]


def _iter_source_files(roots: list[Path]) -> list[Path]:
    found: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        if root.is_file():
            found.append(root)
            continue
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            if any(part in SKIP_DIR_NAMES for part in path.parts):
                continue
            found.append(path)
    return sorted(found)


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _extract_table(path: Path) -> dict[str, Any]:
    import pandas as pd

    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path)
        sheets = {"csv": _frame_preview(frame)}
    else:
        sheets = {
            str(name): _frame_preview(frame)
            for name, frame in pd.read_excel(path, sheet_name=None).items()
        }
    return {"kind": "table", "sheets": sheets, "sheet_count": len(sheets)}


def _frame_preview(frame: Any) -> dict[str, Any]:
    columns = [str(col) for col in list(frame.columns)]
    row_count = int(len(frame))
    preview = frame.head(20).fillna("").astype(str).to_dict(orient="records")
    return {
        "columns": columns,
        "row_count": row_count,
        "preview_rows": preview,
    }


def _extract_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(path))
    paragraphs = [para.text for para in document.paragraphs if para.text.strip()]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    return {
        "kind": "docx",
        "paragraphs": paragraphs[:200],
        "table_count": len(tables),
        "tables": tables[:20],
        "text_preview": "\n".join(paragraphs)[:TEXT_PREVIEW_CHARS],
    }


def _extract_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "has_text_frame") and shape.has_text_frame
        ]
        slides.append({"index": index, "text": "\n".join(texts)})
    return {"kind": "pptx", "slide_count": len(slides), "slides": slides}


def _pdf_has_text(path: Path) -> tuple[bool, str, int]:
    import pdfplumber

    chunks: list[str] = []
    page_count = 0
    with pdfplumber.open(str(path)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(text)
    joined = "\n".join(chunks)
    return len(joined.strip()) >= PDF_TEXT_MIN_CHARS, joined[:TEXT_PREVIEW_CHARS], page_count


def _extract_digital_pdf(path: Path, text_preview: str, page_count: int) -> dict[str, Any]:
    import pdfplumber

    tables: list[list[list[str | None]]] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            tables.extend(page.extract_tables() or [])
    return {
        "kind": "pdf_text",
        "page_count": page_count,
        "table_count": len(tables),
        "tables": tables[:30],
        "text_preview": text_preview,
    }


def _rasterize_pdf(path: Path, pages_dir: Path, file_id: str) -> list[str]:
    pages_dir.mkdir(parents=True, exist_ok=True)
    prefix = pages_dir / f"{file_id}-p"
    try:
        subprocess.run(
            ["pdftoppm", "-png", str(path), str(prefix)],
            check=True,
            capture_output=True,
            text=True,
        )
    except (OSError, subprocess.CalledProcessError) as exc:
        raise RuntimeError(f"pdftoppm failed: {exc}") from exc
    return sorted(str(item) for item in pages_dir.glob(f"{file_id}-p*.png"))


def _ocr_image(path: Path) -> str | None:
    try:
        completed = subprocess.run(
            ["tesseract", str(path), "stdout"],
            check=True,
            capture_output=True,
            text=True,
        )
    except (OSError, subprocess.CalledProcessError):
        return None
    text = completed.stdout.strip()
    return text or None


def _extract_image(path: Path) -> dict[str, Any]:
    ocr = _ocr_image(path)
    return {
        "kind": "image",
        "ocr_text": ocr,
        "needs_vision": ocr is None,
        "text_preview": (ocr or "")[:TEXT_PREVIEW_CHARS],
    }


def ingest_file(path: Path, out_dir: Path) -> ManifestEntry:
    parser = classify_parser(path)
    file_id = file_id_for(path)
    extracted_path = out_dir / "extracted" / f"{file_id}.json"
    rel = str(path)
    if parser == "skip":
        return ManifestEntry(
            file_id=file_id,
            path=rel,
            parser=parser,
            status="skipped",
            confidence=0.0,
            error="unsupported type",
        )
    try:
        payload: dict[str, Any]
        confidence = 0.9
        needs_vision = False
        page_count = None
        sheet_count = None
        if parser == "pandas":
            payload = _extract_table(path)
            sheet_count = int(payload.get("sheet_count") or 0)
        elif parser == "python-docx":
            payload = _extract_docx(path)
        elif parser == "python-pptx":
            payload = _extract_pptx(path)
            page_count = int(payload.get("slide_count") or 0)
        elif parser == "pdfplumber":
            has_text, preview, page_count = _pdf_has_text(path)
            if has_text:
                payload = _extract_digital_pdf(path, preview, page_count)
            else:
                page_paths = _rasterize_pdf(
                    path, out_dir / "ingest" / "pages", file_id
                )
                payload = {
                    "kind": "pdf_scan",
                    "page_count": page_count,
                    "page_images": page_paths,
                    "vision_batch_size": VISION_BATCH_PAGES,
                    "needs_vision": True,
                }
                needs_vision = True
                confidence = 0.4
                parser = "pdftoppm"
        else:
            payload = _extract_image(path)
            needs_vision = bool(payload.get("needs_vision"))
            confidence = 0.6 if payload.get("ocr_text") else 0.3
        payload["source_path"] = rel
        payload["file_id"] = file_id
        _write_json(extracted_path, payload)
        return ManifestEntry(
            file_id=file_id,
            path=rel,
            parser=parser,
            status="needs_vision" if needs_vision else "ok",
            confidence=confidence,
            extracted_path=str(extracted_path),
            page_count=page_count,
            sheet_count=sheet_count,
            needs_vision=needs_vision,
        )
    except Exception as exc:
        return ManifestEntry(
            file_id=file_id,
            path=rel,
            parser=parser,
            status="error",
            confidence=0.0,
            error=str(exc)[:500],
        )


def write_exceptions(out_dir: Path, entries: list[ManifestEntry]) -> Path | None:
    failed = [entry for entry in entries if entry.status in {"error", "skipped"}]
    if not failed:
        return None
    path = out_dir / "exceptions" / "ingest.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=["file_id", "path", "parser", "status", "error"]
        )
        writer.writeheader()
        for entry in failed:
            writer.writerow(
                {
                    "file_id": entry.file_id,
                    "path": entry.path,
                    "parser": entry.parser,
                    "status": entry.status,
                    "error": entry.error or "",
                }
            )
    return path


def run_ingest(
    *,
    roots: list[Path],
    out_dir: Path,
    limit: int | None = None,
    dry_run: bool = False,
) -> IngestResult:
    result = IngestResult()
    files = _iter_source_files(roots)
    if limit is not None:
        files = files[:limit]
    for path in files:
        if dry_run:
            result.entries.append(
                ManifestEntry(
                    file_id=file_id_for(path),
                    path=str(path),
                    parser=classify_parser(path),
                    status="planned",
                    confidence=0.0,
                )
            )
            continue
        result.entries.append(ingest_file(path, out_dir))
    if not dry_run:
        _write_json(out_dir / "ingest" / "MANIFEST.json", result.to_manifest())
        write_exceptions(out_dir, result.entries)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Ingest Craft source files")
    parser.add_argument(
        "--roots",
        nargs="+",
        default=["project", "user_library", "attachments"],
    )
    parser.add_argument("--out", default="outputs")
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    result = run_ingest(
        roots=[Path(root) for root in args.roots],
        out_dir=Path(args.out),
        limit=args.limit,
        dry_run=args.dry_run,
    )
    print(json.dumps(result.to_manifest(), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
